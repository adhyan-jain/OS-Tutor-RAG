"""Evaluate pipeline configurations against eval_set.json using RAGAS metrics.

Runs RAGPipeline across a set of technique combinations (config variants),
scores each with RAGAS (see src.evaluation.METRIC_NAMES for the full metric
list) using a local judge (Ollama LLM + a local sentence-transformer for
embeddings, not OpenAI), and writes a comparison table ranking configs by
correctness + completeness.

This is the *mass* evaluator (every technique combination, one full run).
For evaluating a single pipeline config as you iterate, use
RAGPipeline.evaluate() instead (src/pipeline.py), which appends to a shared
per-run workbook via src/evaluation.py -- both share the same scoring logic.
"""

from __future__ import annotations

import json
import math
from dataclasses import dataclass
from pathlib import Path

from src.config import GenerationConfig, PathConfig, PipelineConfig
from src.evaluation import (
    METRIC_NAMES,
    append_run_to_workbook,
    composite_scores,
    load_eval_set,
    score_pipeline,
)
from src.ingestion.extract_docx import extract_docx
from src.ingestion.extract_pdf import extract_pdf
from src.ingestion.extract_ppt import extract_ppt
from src.pipeline import RAGPipeline

_EXTRACTORS = {
    ".pptx": lambda path: [extract_ppt(path)],
    ".docx": lambda path: extract_docx(path),
    ".pdf": lambda path: [extract_pdf(path)],
}


def load_raw_documents(raw_dir: Path):
    """Extract Documents from every supported file in data/raw/.

    Args:
        raw_dir: Directory to scan (e.g. PathConfig().data_raw_dir).

    Returns:
        A flat list of Documents across all supported raw files.
    """
    documents = []
    for path in sorted(raw_dir.iterdir()):
        extractor = _EXTRACTORS.get(path.suffix.lower())
        if extractor is not None:
            documents.extend(extractor(path))
    return documents


@dataclass
class ConfigVariant:
    """A named PipelineConfig variant to evaluate."""

    name: str
    config: PipelineConfig


def default_config_variants(generation_config: GenerationConfig | None = None) -> list[ConfigVariant]:
    """Build the grid of config variants to compare: retrieval technique x
    use_multi_query (where valid) x reranking method x MMR on/off.

    Chunking is deliberately not an axis. Each source type is matched to the
    strategy that suits how that format is written (see ChunkingConfig), so
    forcing one strategy across every type compares formats rather than
    strategies -- semantic chunking a slide deck, for instance, collapses it
    into a handful of oversized chunks because slide bullets carry no sentence
    punctuation to split on.

    use_multi_query is excluded for technique == "hyde" (see
    RetrievalConfig.use_multi_query's docstring -- the two aren't combinable).

    "llm_rerank" is deliberately excluded despite being a supported
    RerankingConfig.method: it prompts the LLM once per retrieved candidate,
    which roughly triples the sweep for a stage cross_encoder already covers.
    Evaluate it separately against whichever config wins.

    Args:
        generation_config: Generation settings shared by all variants.
            Defaults to GenerationConfig() (currently backend="ollama",
            model_name="llama3:latest" -- see config.py).

    Returns:
        A list of named ConfigVariants (28 by default: 7 retrieval configs x
        2 reranking methods x 2 MMR settings).
    """
    generation_config = generation_config or GenerationConfig()

    variants = []
    for technique in ("dense", "bm25", "hybrid_rrf", "hyde"):
        multi_query_options = (False,) if technique == "hyde" else (False, True)
        for use_multi_query in multi_query_options:
            for method in ("cross_encoder", "none"):
                for mmr_enabled in (True, False):
                    config = PipelineConfig(generation=generation_config)
                    config.retrieval.technique = technique
                    config.retrieval.use_multi_query = use_multi_query
                    config.reranking.method = method
                    config.diversification.enabled = mmr_enabled

                    retrieval_label = f"{technique}+multi_query" if use_multi_query else technique
                    name = f"{retrieval_label}+{method}+mmr_{mmr_enabled}"
                    variants.append(ConfigVariant(name=name, config=config))
    return variants


def focused_config_variants(generation_config: GenerationConfig | None = None) -> list[ConfigVariant]:
    """The configurations still worth spending RAGAS time on.

    Retrieval was characterised first, without any LLM in the loop
    (eval/retrieval_eval.py), which narrowed six techniques to three: dense
    leads MRR and recall@5, HyDE leads full@5, and multi-query leads recall@1,
    each for a mechanistic reason (FINDINGS.md A12). bm25 and hybrid_rrf are
    dominated on every retrieval metric, so their answer quality is not in
    question and full-sweep time spent on them buys documentation, not
    findings.

    HyDE and multi-query are paired as well: their advantages are complementary
    (breadth against precision), and the published HyDE samples several
    hypothetical documents rather than one, so the combination is closer to the
    canonical method than the single-shot version.

    Reranking stays as an axis despite measuring harmful (A7), because the
    negative result is worth having on the corrected benchmark rather than only
    on the one A11 showed was biased. Diversification does not: A10 measured it
    dropping correct slides outright.
    """
    generation_config = generation_config or GenerationConfig()

    retrieval_options = [
        ("dense", "dense", False),
        ("hyde", "hyde", False),
        ("dense+multi_query", "dense", True),
        ("hyde+multi_query", "hyde", True),
    ]

    variants = []
    for label, technique, use_multi_query in retrieval_options:
        for method in ("none", "cross_encoder"):
            config = PipelineConfig(generation=generation_config)
            config.retrieval.technique = technique
            config.retrieval.use_multi_query = use_multi_query
            config.reranking.method = method
            config.diversification.enabled = False
            variants.append(ConfigVariant(name=f"{label}+{method}", config=config))
    return variants


def prompt_config_variants(styles: list[str] | None = None) -> list[ConfigVariant]:
    """The answer prompts, held against one fixed retrieval configuration.

    A17 measured the prompt worth more than every retrieval technique put
    together, so prompts get their own axis. Everything downstream of retrieval
    is pinned to the current default winner (dense + multi_query with five
    variants, cross-encoder reranking, no diversification) precisely so the
    only difference between rows is the wording of the prompt.

    "cot" gets a larger max_tokens: it spends output on a reasoning section
    before the answer, and holding the budget at the default would confound
    reasoning's effect with truncation of the answer that follows it.
    """
    styles = styles or ["few_shot", "cot"]

    variants = []
    for style in styles:
        generation_config = GenerationConfig(prompt_style=style)
        if style in ("cot",):
            generation_config.max_tokens = 768
        config = PipelineConfig(generation=generation_config)
        config.retrieval.technique = "dense"
        config.retrieval.use_multi_query = True
        config.reranking.method = "cross_encoder"
        config.diversification.enabled = False
        variants.append(ConfigVariant(name=f"prompt_{style}", config=config))
    return variants


def completed_run_names(workbook_path: Path) -> set[str]:
    """Names of variants already recorded in the per-run workbook.

    Used to resume a sweep that died partway: a full-grid run takes many hours,
    and re-running variants that already produced results wastes most of that.
    Reads the "Final Analysis" sheet, whose first column is the run name.

    Args:
        workbook_path: Path to the per-run workbook (may not exist yet).

    Returns:
        The set of run names present, empty if the workbook has no results.
    """
    if not workbook_path.exists():
        return set()

    import openpyxl

    workbook = openpyxl.load_workbook(workbook_path, read_only=True)
    if "Final Analysis" not in workbook.sheetnames:
        return set()

    sheet = workbook["Final Analysis"]
    names = {
        str(row[0]) for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]
    }
    workbook.close()
    return names


def run_comparison(
    variants: list[ConfigVariant],
    documents,
    eval_set_path: Path,
    output_path: Path,
    details_output_path: Path | None = None,
    resume: bool = False,
) -> None:
    """Run RAGAS eval across config variants and write a ranked comparison table.

    Also writes a per-question detail workbook (one sheet per config variant,
    plus an "All Results (by question)" sheet grouping every variant's answer
    for the same question in adjacent rows) so answers aren't just summarized
    away into the aggregate table.

    Args:
        variants: Named PipelineConfig variants to evaluate.
        documents: Documents to ingest into each variant's pipeline before eval.
        eval_set_path: Path to the eval_set.json file.
        output_path: Where to write the comparison table (.md or .csv, by extension).
        details_output_path: Where to write the per-question Excel workbook.
            Defaults to output_path's directory / "ragas_details.xlsx".
        resume: Skip variants already recorded in the per-run workbook, for
            restarting a sweep that died partway through.
    """
    eval_set = load_eval_set(eval_set_path)
    rows = []
    all_detail_rows: list[dict] = []
    details_output_path = details_output_path or output_path.parent / "ragas_details.xlsx"

    if resume:
        already_done = completed_run_names(variants[0].config.eval.output_workbook_path)
        skipped = [v.name for v in variants if v.name in already_done]
        variants = [v for v in variants if v.name not in already_done]
        print(f"Resuming: skipping {len(skipped)} already-completed variant(s), {len(variants)} remaining\n")

    failures: list[tuple[str, str]] = []

    for variant_index, variant in enumerate(variants, start=1):
        print(f"[{variant_index}/{len(variants)}] Evaluating variant: {variant.name}")

        # One variant failing (a flaky Ollama call, an OOM, a metric blowing
        # up on odd output) must not discard the hours of completed variants
        # before it -- log it, keep the partial results already written, and
        # continue with the next variant.
        try:
            pipeline = RAGPipeline(variant.config)
            pipeline.ingest(documents)
            scores, detail_rows, eval_tracker = score_pipeline(pipeline, eval_set)
        except Exception as error:  # noqa: BLE001 - deliberate: keep sweep alive
            print(f"  FAILED ({type(error).__name__}): {error}")
            failures.append((variant.name, f"{type(error).__name__}: {error}"))
            continue

        rows.append({"config": variant.name, **scores})
        for question_index, detail_row in enumerate(detail_rows):
            all_detail_rows.append({"config": variant.name, "question_index": question_index, **detail_row})
        print(f"  {scores}")

        # Also record into the shared per-run workbook (same LLM usage/cost
        # reporting as RAGPipeline.evaluate()), so mass sweeps and single
        # ad-hoc runs both accumulate into one comparable history.
        append_run_to_workbook(
            variant.config.eval.output_workbook_path,
            variant.name,
            variant.config,
            scores,
            detail_rows,
            pipeline.tracker,
            eval_tracker,
        )

        # Write after every variant (not just at the end) so a crash midway
        # through this multi-hour run doesn't lose already-computed results.
        _write_details_workbook(all_detail_rows, details_output_path)
        _write_table(_ranked(rows), output_path)

    if failures:
        print(f"\n{len(failures)} variant(s) failed and were skipped:")
        for name, error in failures:
            print(f"  {name}: {error}")

    print(f"Comparison table written to {output_path}")
    print(f"Per-question details written to {details_output_path}")


def _ranked(rows: list[dict]) -> list[dict]:
    """Rank configs by correctness + completeness, not a flat metric average.

    Composites come from src.evaluation.composite_scores so this comparison
    table and the per-run workbook rank runs by identical definitions -- they
    previously diverged (completeness was context_recall alone here, but
    mean(context_recall, context_entity_recall) in the workbook).

    Sorting is NaN-safe: a config whose composite failed to compute sorts last
    rather than comparing unpredictably against real scores.
    """
    ranked_rows = []
    for row in rows:
        ranked_row = dict(row)
        ranked_row.update(composite_scores(row))
        ranked_rows.append(ranked_row)
    ranked_rows.sort(
        key=lambda r: (
            -1.0 if math.isnan(r["correctness_completeness"]) else r["correctness_completeness"]
        ),
        reverse=True,
    )
    return ranked_rows


def _variant_sheet_names(variant_names: list[str]) -> dict[str, str]:
    """Map each variant name to a unique, Excel-legal (<=31 char) sheet name.

    Variant names are far longer than Excel's 31-character sheet-name limit
    (e.g. "structure_aware+dense+cross_encoder+mmr_True"), and truncating them
    collides -- variants differing only in a trailing field would map to the
    same sheet. A numeric prefix keeps every sheet distinct regardless of how
    much of the name survives truncation; the untruncated name stays available
    in each sheet's "config" column and in the summary table.
    """
    return {
        name: f"{index:02d}_{name.replace('+', '_')}"[:31]
        for index, name in enumerate(variant_names, start=1)
    }


def _write_details_workbook(all_detail_rows: list[dict], output_path: Path) -> None:
    """Write per-question results in a side-by-side-comparable layout.

    The "All Results (by question)" sheet is sorted by question first, then
    config, so every variant's answer/scores for the same question sit in
    adjacent rows -- easy to scan across configs for one question at a time,
    rather than having to flip between per-variant sheets.
    """
    import pandas as pd

    df = pd.DataFrame(all_detail_rows)
    by_question = df.sort_values(["question_index", "config"]).drop(columns=["question_index"])

    variant_names = list(df["config"].unique())
    sheet_names = _variant_sheet_names(variant_names)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        by_question.to_excel(writer, sheet_name="All Results (by question)", index=False)
        for variant_name in variant_names:
            variant_df = df[df["config"] == variant_name].drop(columns=["question_index"])
            variant_df.to_excel(writer, sheet_name=sheet_names[variant_name], index=False)


def _write_table(rows: list[dict], output_path: Path) -> None:
    columns = ["config", *METRIC_NAMES, "correctness", "completeness", "correctness_completeness"]

    if output_path.suffix == ".csv":
        import csv

        with open(output_path, "w", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)
        return

    lines = [
        "| " + " | ".join(columns) + " |",
        "| " + " | ".join(["---"] * len(columns)) + " |",
    ]
    for row in rows:
        lines.append(
            "| " + " | ".join(f"{row[c]:.4f}" if c != "config" else row[c] for c in columns) + " |"
        )
    output_path.write_text("\n".join(lines) + "\n")


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--quick",
        action="store_true",
        help="Smoke-test mode: first 3 eval questions, 2 config variants only.",
    )
    parser.add_argument(
        "--output",
        default="eval/comparison_results.md",
        help="Output path for the comparison table (.md or .csv).",
    )
    parser.add_argument(
        "--resume",
        action="store_true",
        help="Skip variants already recorded in the per-run workbook (restart a died sweep).",
    )
    parser.add_argument(
        "--num-questions",
        type=int,
        default=None,
        help="Evaluate only the first N eval questions (default: the whole eval set).",
    )
    parser.add_argument(
        "--focused",
        action="store_true",
        help="Only the configs retrieval measurement left undominated (see focused_config_variants).",
    )
    parser.add_argument(
        "--prompts",
        nargs="*",
        metavar="STYLE",
        help="Compare answer prompts on one fixed retrieval config "
        "(default: few_shot cot). See prompt_config_variants.",
    )
    args = parser.parse_args()

    documents = load_raw_documents(PathConfig().data_raw_dir)
    print(f"Loaded {len(documents)} documents from data/raw/")

    if args.prompts is not None:
        variants = prompt_config_variants(args.prompts or None)
    elif args.focused:
        variants = focused_config_variants()
    else:
        variants = default_config_variants()
    eval_set_path = Path("eval/eval_set.json")

    if args.quick:
        variants = variants[:2]
        eval_set = load_eval_set(eval_set_path)[:3]
        quick_eval_set_path = Path("eval/_quick_eval_set.json")
        quick_eval_set_path.write_text(json.dumps({"examples": eval_set}))
        eval_set_path = quick_eval_set_path
    elif args.num_questions is not None:
        eval_set = load_eval_set(eval_set_path)[: args.num_questions]
        subset_path = Path("eval/_subset_eval_set.json")
        subset_path.write_text(json.dumps({"examples": eval_set}))
        eval_set_path = subset_path

    print(f"Running {len(variants)} config variants")
    run_comparison(variants, documents, eval_set_path, Path(args.output), resume=args.resume)


if __name__ == "__main__":
    main()
